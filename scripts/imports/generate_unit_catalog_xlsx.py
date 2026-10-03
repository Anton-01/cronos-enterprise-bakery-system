#!/usr/bin/env python3
"""Generates the unit-catalog .xlsx files consumed by the bulk import endpoints.

Outputs (paths relative to the repository root):
  src/main/resources/imports/templates/unit-types-template.xlsx         empty template (GET /unit-type/import/template)
  src/main/resources/imports/templates/measurement-units-template.xlsx  empty template (GET /measurement-unit/import/template)
  docs/imports/01-unit-types.xlsx                                       full bakery catalog, import FIRST
  docs/imports/02-measurement-units.xlsx                                full bakery catalog, import SECOND

The data sheets must keep the exact names "UnitTypes" / "MeasurementUnits" and the headers in row 1;
every other sheet (instructions) is ignored by the importer. Re-run after editing the catalog below:

    pip install openpyxl==3.1.5 && python3 scripts/imports/generate_unit_catalog_xlsx.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
TEMPLATES = ROOT / "src/main/resources/imports/templates"
DATA = ROOT / "docs/imports"

DIMENSIONS = ["MASS", "VOLUME", "COUNT", "LENGTH"]
STATUSES = ["ACTIVE", "INACTIVE", "ARCHIVED"]
MAX_ROWS = 2000

UNIT_TYPE_COLUMNS = [  # (header, width, description)
    ("code", 14, "Obligatorio. Identificador único (sin importar mayúsculas). Letras, dígitos, '.', '_' o '-'; máx. 20."),
    ("name", 26, "Obligatorio. Nombre visible, único. Máx. 100 caracteres."),
    ("dimension", 14, "Obligatorio. MASS | VOLUME | COUNT | LENGTH. Solo un tipo de unidad por dimensión."),
    ("status", 12, "Opcional. ACTIVE | INACTIVE | ARCHIVED. Vacío = conserva el actual (ACTIVE si es nuevo)."),
]

MEASUREMENT_UNIT_COLUMNS = [
    ("code", 12, "Obligatorio. Identificador único y SENSIBLE a mayúsculas ('T' ≠ 't'). Sin espacios; máx. 20."),
    ("name", 24, "Obligatorio. Nombre en singular, único. Máx. 100."),
    ("namePlural", 24, "Obligatorio. Nombre en plural. Máx. 100."),
    ("unitTypeCode", 14, "Obligatorio. 'code' de un tipo de unidad existente (importa primero los tipos de unidad)."),
    ("factorToBase", 16, "Obligatorio. Cuántas unidades base equivalen a 1 de esta unidad (kg = 1000 si la base es g). > 0, máx. 10 decimales, punto decimal."),
    ("isBaseUnit", 12, "Obligatorio. TRUE | FALSE. Exactamente una unidad base por tipo, con factor 1."),
    ("status", 12, "Opcional. ACTIVE | INACTIVE | ARCHIVED. Vacío = conserva el actual (ACTIVE si es nuevo)."),
]

UNIT_TYPES = [
    ("MASS", "Masa", "MASS", "ACTIVE"),
    ("VOLUME", "Volumen", "VOLUME", "ACTIVE"),
    ("COUNT", "Conteo", "COUNT", "ACTIVE"),
    ("LENGTH", "Longitud", "LENGTH", "ACTIVE"),
]

# Volume comes in two internally consistent families:
#  - kitchen measures (recipes): the US FDA "legal" household measures used on nutrition labels
#    (21 CFR 101.9(b)(5)(viii)): 1 cup = 240 mL, 1 tbsp = 15 mL, 1 tsp = 5 mL (1 cup = 16 tbsp = 48 tsp),
#    plus the 250 mL metric cup;
#  - US customary liquid measures (purchasing): exact by definition, 1 gal = 231 in³ = 3785.411784 mL
#    = 4 qt = 8 pt = 128 fl oz.
# The legal cup (240 mL) is deliberately NOT 8 customary fl oz (236.59 mL).
# Mass: avoirdupois ounce/pound are exact (1 lb = 453.59237 g, 1 oz = 1/16 lb).
MEASUREMENT_UNITS = [
    # MASS — base: gram
    ("g", "gramo", "gramos", "MASS", 1, True),
    ("mg", "miligramo", "miligramos", "MASS", 0.001, False),
    ("kg", "kilogramo", "kilogramos", "MASS", 1000, False),
    ("oz", "onza", "onzas", "MASS", 28.349523125, False),
    ("lb", "libra", "libras", "MASS", 453.59237, False),
    # VOLUME — base: millilitre
    ("ml", "mililitro", "mililitros", "VOLUME", 1, True),
    ("dl", "decilitro", "decilitros", "VOLUME", 100, False),
    ("l", "litro", "litros", "VOLUME", 1000, False),
    ("tsp", "cucharadita", "cucharaditas", "VOLUME", 5, False),
    ("tbsp", "cucharada", "cucharadas", "VOLUME", 15, False),
    ("fl_oz", "onza líquida (EUA)", "onzas líquidas (EUA)", "VOLUME", 29.5735295625, False),
    ("cup", "taza", "tazas", "VOLUME", 240, False),
    ("cup_m", "taza métrica", "tazas métricas", "VOLUME", 250, False),
    ("pt", "pinta", "pintas", "VOLUME", 473.176473, False),
    ("qt", "cuarto de galón", "cuartos de galón", "VOLUME", 946.352946, False),
    ("gal", "galón", "galones", "VOLUME", 3785.411784, False),
    # COUNT — base: piece
    ("pz", "pieza", "piezas", "COUNT", 1, True),
    ("par", "par", "pares", "COUNT", 2, False),
    ("mdz", "media docena", "medias docenas", "COUNT", 6, False),
    ("dz", "docena", "docenas", "COUNT", 12, False),
    ("cto", "ciento", "cientos", "COUNT", 100, False),
    # LENGTH — base: centimetre (molds, trays, cake diameters)
    ("cm", "centímetro", "centímetros", "LENGTH", 1, True),
    ("mm", "milímetro", "milímetros", "LENGTH", 0.1, False),
    ("m", "metro", "metros", "LENGTH", 100, False),
    ("in", "pulgada", "pulgadas", "LENGTH", 2.54, False),
]

HEADER_FILL = PatternFill("solid", fgColor="1F7A5C")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TEXT_FORMAT = "@"


def data_sheet(wb, title, columns, rows, validations):
    ws = wb.active
    ws.title = title
    for idx, (header, width, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"

    for r, values in enumerate(rows, start=2):
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value)

    # Text columns stay text even when someone types "1" or "0.5" as a code.
    text_columns = [i for i, (h, _, _) in enumerate(columns, start=1) if h in ("code", "name", "namePlural", "unitTypeCode")]
    for col in text_columns:
        for r in range(2, MAX_ROWS + 2):
            ws.cell(row=r, column=col).number_format = TEXT_FORMAT

    for header, rule in validations.items():
        col = get_column_letter([h for h, _, _ in columns].index(header) + 1)
        rule.add(f"{col}2:{col}{MAX_ROWS + 1}")
        ws.add_data_validation(rule)
    return ws


def instructions_sheet(wb, title, columns, notes):
    ws = wb.create_sheet("Instrucciones")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"], ws["B3"] = "Columna", "Regla"
    ws["A3"].font = ws["B3"].font = Font(bold=True)
    for i, (header, _, description) in enumerate(columns, start=4):
        ws.cell(row=i, column=1, value=header)
        ws.cell(row=i, column=2, value=description).alignment = Alignment(wrap_text=True)
    start = 5 + len(columns)
    for i, note in enumerate(notes, start=start):
        ws.cell(row=i, column=1, value="•")
        ws.cell(row=i, column=2, value=note).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 110


def list_rule(values, title):
    rule = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    rule.error, rule.errorTitle = f"Valores permitidos: {', '.join(values)}", title
    rule.showErrorMessage = True
    return rule


def factor_rule():
    rule = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True)
    rule.error, rule.errorTitle = "Debe ser un número mayor que cero (punto decimal)", "factorToBase"
    rule.showErrorMessage = True
    return rule


COMMON_NOTES = [
    "La hoja de datos debe conservar su nombre exacto y los encabezados en la fila 1; esta hoja se ignora.",
    "La importación es todo-o-nada: si una sola fila tiene error, no se guarda nada y el reporte indica fila, columna y motivo.",
    "Primero sube el archivo en modo validación (dryRun=true); si el reporte sale VALIDATED, súbelo en modo aplicar (dryRun=false).",
    "Las filas se identifican por 'code': si el código existe se actualiza, si no existe se crea. Nada se elimina desde el archivo.",
    "Límites: 2 MB, 2000 filas, solo .xlsx sin macros ni contraseña. No uses fórmulas con error (#N/A, #DIV/0!).",
]

UNIT_NOTES = COMMON_NOTES + [
    "Importa SIEMPRE primero los tipos de unidad (01-unit-types.xlsx) y después las unidades de medida.",
    "Unidades en uso por insumos o recetas: su tipo, factor y marca de base no pueden cambiar (protege los costos históricos).",
    "Los códigos g, cup, tbsp y tsp están reservados por el sistema (densidades de ingredientes): no pueden renombrarse ni desactivarse.",
    "cup/tbsp/tsp siguen las medidas 'legales' de etiquetado nutrimental (1 taza = 240 mL, 1 cda = 15 mL, 1 cdita = 5 mL); fl_oz/pt/qt/gal son medidas usuales de EUA (1 gal = 128 fl oz = 3785.411784 mL).",
]


def build(path, title, sheet_name, columns, rows, validations, notes):
    wb = Workbook()
    data_sheet(wb, sheet_name, columns, rows, validations)
    instructions_sheet(wb, title, columns, notes)
    wb.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"wrote {path.relative_to(ROOT)} ({len(rows)} rows)")


def unit_type_validations():
    return {"dimension": list_rule(DIMENSIONS, "dimension"), "status": list_rule(STATUSES, "status")}


def measurement_unit_validations():
    return {"isBaseUnit": list_rule(["TRUE", "FALSE"], "isBaseUnit"), "status": list_rule(STATUSES, "status"),
            "factorToBase": factor_rule()}


def main():
    unit_rows = [(c, n, p, t, f, b, "ACTIVE") for c, n, p, t, f, b in MEASUREMENT_UNITS]
    build(TEMPLATES / "unit-types-template.xlsx", "Plantilla de tipos de unidad", "UnitTypes",
          UNIT_TYPE_COLUMNS, [], unit_type_validations(), COMMON_NOTES)
    build(TEMPLATES / "measurement-units-template.xlsx", "Plantilla de unidades de medida", "MeasurementUnits",
          MEASUREMENT_UNIT_COLUMNS, [], measurement_unit_validations(), UNIT_NOTES)
    build(DATA / "01-unit-types.xlsx", "Catálogo de tipos de unidad (importar primero)", "UnitTypes",
          UNIT_TYPE_COLUMNS, UNIT_TYPES, unit_type_validations(), COMMON_NOTES)
    build(DATA / "02-measurement-units.xlsx", "Catálogo de unidades de medida (importar segundo)", "MeasurementUnits",
          MEASUREMENT_UNIT_COLUMNS, unit_rows, measurement_unit_validations(), UNIT_NOTES)


if __name__ == "__main__":
    main()
